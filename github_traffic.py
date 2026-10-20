"""
GitHub Traffic Report
Holt für alle eigenen Repos: Aufrufe, Besucher, Clones,
Herkunftsseiten und geöffnete Dateien.
Speichert den Verlauf als CSV und erstellt einen Excel-Bericht.
"""

import os
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ── 1. Einstellungen ──────────────────────────────────────────
TOKEN = os.environ.get("GITHUB_TOKEN")
if not TOKEN:
    sys.exit("GITHUB_TOKEN fehlt. Bitte als Umgebungsvariable setzen.")

API = "https://api.github.com"
HEADERS = {
    "Authorization": f"Bearer {TOKEN}",
    "Accept": "application/vnd.github+json",
    "X-GitHub-Api-Version": "2022-11-28",
}
OUT_DIR = Path(__file__).parent
OUT_DIR.mkdir(exist_ok=True)
TODAY = date.today().isoformat()


# ── 2. Hilfsfunktionen ────────────────────────────────────────
def get(url, params=None):
    """GET-Anfrage mit Fehlerbehandlung."""
    r = requests.get(url, headers=HEADERS, params=params, timeout=30)
    if r.status_code == 401:
        sys.exit("Token ungültig oder abgelaufen.")
    r.raise_for_status()
    return r.json()


def get_repos():
    """Alle eigenen Repos holen."""
    repos, page = [], 1
    while True:
        batch = get(f"{API}/user/repos",
                    {"affiliation": "owner", "per_page": 100, "page": page})
        if not batch:
            break
        repos += batch
        page += 1
    return repos


def save_history(df, filename, keys):
    """Neue Daten an CSV anhängen, Duplikate entfernen."""
    path = OUT_DIR / filename
    if path.exists():
        df = pd.concat([pd.read_csv(path), df], ignore_index=True)
    df = df.drop_duplicates(subset=keys, keep="last")
    df.to_csv(path, index=False)


def load(filename, cols):
    """Gespeicherten Verlauf laden."""
    p = OUT_DIR / filename
    return pd.read_csv(p) if p.exists() else pd.DataFrame(columns=cols)


def style(ws, highlight_col=None):
    """Excel-Blatt formatieren."""
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if highlight_col:
        green = PatternFill("solid", fgColor="C6EFCE")
        for row in ws.iter_rows(min_row=2):
            if (row[highlight_col - 1].value or 0) > 0:
                for c in row:
                    c.fill = green


# ── 3. Daten abrufen ──────────────────────────────────────────
daily, referrers, paths, summary = [], [], [], []

for repo in get_repos():
    name, full = repo["name"], repo["full_name"]
    try:
        views = get(f"{API}/repos/{full}/traffic/views")
        clones = get(f"{API}/repos/{full}/traffic/clones")
        refs = get(f"{API}/repos/{full}/traffic/popular/referrers")
        pops = get(f"{API}/repos/{full}/traffic/popular/paths")
    except requests.HTTPError as e:
        print(f"Übersprungen: {name} ({e})")
        continue

    for v in views.get("views", []):
        daily.append({"repo": name, "date": v["timestamp"][:10],
                      "views": v["count"], "unique_visitors": v["uniques"]})
    for r in refs:
        referrers.append({"repo": name, "snapshot": TODAY, "referrer": r["referrer"],
                          "views": r["count"], "unique": r["uniques"]})
    for p in pops:
        paths.append({"repo": name, "snapshot": TODAY, "path": p["path"],
                      "views": p["count"], "unique": p["uniques"]})
    summary.append({
        "repo": name,
        "views_14d": views.get("count", 0),
        "unique_14d": views.get("uniques", 0),
        "clones_14d": clones.get("count", 0),
        "stars": repo["stargazers_count"],
        "forks": repo["forks_count"],
        "top_referrer": refs[0]["referrer"] if refs else "",
    })


# ── 4. Verlauf als CSV speichern ──────────────────────────────
if daily:
    save_history(pd.DataFrame(daily), "views_daily.csv", ["repo", "date"])
if referrers:
    save_history(pd.DataFrame(referrers), "referrers.csv", ["repo", "snapshot", "referrer"])
if paths:
    save_history(pd.DataFrame(paths), "paths.csv", ["repo", "snapshot", "path"])


# ── 5. Tabellen für Excel vorbereiten ─────────────────────────
uebersicht = pd.DataFrame(summary).sort_values("views_14d", ascending=False)
uebersicht.columns = ["Projekt", "Aufrufe", "Besucher", "Clones",
                      "Stars", "Forks", "Hauptquelle"]

verlauf = load("views_daily.csv", ["repo", "date", "views", "unique_visitors"])
verlauf = verlauf[verlauf["views"] > 0].sort_values(["date", "repo"], ascending=[False, True])
verlauf.columns = ["Projekt", "Datum", "Aufrufe", "Besucher"]

herkunft = load("referrers.csv", ["repo", "snapshot", "referrer", "views", "unique"])
herkunft = herkunft[herkunft["snapshot"] == TODAY].drop(columns="snapshot")
herkunft.columns = ["Projekt", "Quelle", "Aufrufe", "Besucher"]
herkunft = herkunft.sort_values("Aufrufe", ascending=False)

dateien = load("paths.csv", ["repo", "snapshot", "path", "views", "unique"])
dateien = dateien[dateien["snapshot"] == TODAY].drop(columns="snapshot")
dateien.columns = ["Projekt", "Datei", "Aufrufe", "Besucher"]
dateien = dateien.sort_values("Aufrufe", ascending=False)


# ── 6. Excel schreiben ────────────────────────────────────────
xlsx = OUT_DIR / "github_traffic.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
    for sheet, data, hl in [("Übersicht", uebersicht, 2),
                            ("Verlauf", verlauf, None),
                            ("Herkunft", herkunft, None),
                            ("Dateien", dateien, None)]:
        data.to_excel(writer, sheet_name=sheet, index=False)
        style(writer.sheets[sheet], hl)


# ── 7. Kurzausgabe ────────────────────────────────────────────
aktiv = uebersicht[uebersicht["Aufrufe"] > 0]
print(f"\nGitHub {TODAY}: {len(aktiv)} von {len(uebersicht)} Projekten aufgerufen\n")
for _, r in aktiv.iterrows():
    quelle = f", über {r['Hauptquelle']}" if r["Hauptquelle"] else ""
    print(f"• {r['Projekt']}: {r['Aufrufe']} Aufrufe, {r['Besucher']} Besucher{quelle}")
print(f"\nExcel: {xlsx}")
